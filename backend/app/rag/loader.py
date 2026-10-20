"""文档解析模块：按扩展名路由解析器，产出标题结构化的中间表示。

模块职责（design.md 1.1.3-A / 2.1.2）：
- PDF → pypdf 逐页抽取文本；
- DOCX → python-docx 抽段落与表格文本（Heading 样式 → 标题节点）；
- PPTX → python-pptx 按 slide 抽 text_frame（slide 标题 → 标题节点）；
- MD / TXT → 行级文本读取（md 以 ``#`` 行级正则识别标题层级，
  不用 markdown 包 AST——取舍见 design.md 2.9 决策 #14：提示词明确
  「纯文本方式处理 .md 即可」，标题识别只需行级，避免 HTML 转换失真）。

**空文本检测红线**（spec 5.2.1 规则 2，逐字锁定文案）：解析后有效
字符（``str.isalnum``：字母/数字/CJK，排除空白与标点）少于
``MIN_VALID_CHARS = 50`` 时判定为扫描件或图片型文件，抛
``ScanFileError``，由流水线把 ``knowledge_items.status`` 置为
failed、error_msg 写固定文案，**禁止静默标记为 indexed**。

TXT 编码回退链（design.md 2.9 决策 #15）：UTF-8 → GBK → UTF-16
依次严格解码，覆盖中文 Windows 常见来源；全失败抛编码错误。

pypdf / python-docx / python-pptx 做防御性导入（函数内 import），
导入失败按「解析失败」口径抛出（供流水线置 failed，
design.md 2.3.2-B：依赖已验证可安装 ≠ 运行时 import 必然成功）。

本模块全部为同步实现（CPU 段由 services 层 ``asyncio.to_thread``
卸载），不 import api / services（契约 §4 单向依赖）。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import TypedDict

# 空文本检测阈值：有效字符（isalnum）少于该值判定扫描件（spec 5.2.1 规则 2）。
MIN_VALID_CHARS: int = 50

# 扫描件固定文案（spec 逐字锁定，前端按原文展示；为 error_msg 中唯一
# 不带「{环节}失败：」前缀的例外，见 design.md 2.7-B）。
SCAN_FILE_ERROR_MSG: str = "该文件可能是扫描件，未提取到文字（当前版本不支持 OCR）"

# Markdown 行级标题正则：行首 1~6 个 # + 空白 + 标题文本。
_MD_HEADING_RE: re.Pattern[str] = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")

# TXT 解码回退链：依次严格解码，覆盖 UTF-8 / GBK / UTF-16（含 BOM）。
_ENCODING_CHAIN: tuple[str, ...] = ("utf-8", "gbk", "utf-16")


class ScanFileError(Exception):
    """扫描件/图片型文件信号（空文本检测红线触发）。

    Attributes:
        message: 固定文案（SCAN_FILE_ERROR_MSG，spec 逐字锁定）。
    """

    def __init__(self, message: str = SCAN_FILE_ERROR_MSG) -> None:
        """初始化扫描件信号异常。

        Args:
            message: 固定错误文案，默认 SCAN_FILE_ERROR_MSG。
        """
        super().__init__(message)


class Section(TypedDict):
    """解析中间表示的一个标题区间。

    Attributes:
        heading_path: 标题层级路径（如「机器学习核心概念速览 > 六、Transformer
            与注意力机制」）；无标题文档为空串。
        text: 该标题区间下的正文（段落以空行分隔）。
    """

    heading_path: str
    text: str


class ParsedDocument(TypedDict):
    """解析结果中间表示。

    Attributes:
        sections: 标题区间列表（顺序保持原文顺序）。
        full_text: 全文文本（md/txt 为解码原文；pdf/docx/pptx 为抽取
            文本拼接），供空文本检测与调试。
    """

    sections: list[Section]
    full_text: str


# ---------------------------------------------------------------------------
# 编码解码
# ---------------------------------------------------------------------------


def _decode_text_bytes(raw: bytes) -> str:
    """按 UTF-8 → GBK → UTF-16 回退链严格解码文本字节。

    Args:
        raw: 原始字节流。

    Returns:
        str: 解码后的文本。

    Raises:
        ValueError: 三种编码全部失败时抛出（由流水线转
            failed「解析失败：文件编码无法识别」）。
    """
    for encoding in _ENCODING_CHAIN:
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ValueError("文件编码无法识别（已尝试 UTF-8/GBK/UTF-16）")


# ---------------------------------------------------------------------------
# 空文本检测
# ---------------------------------------------------------------------------


def _effective_char_count(text: str) -> int:
    """统计有效字符数（字母/数字/CJK，排除空白与标点）。

    Args:
        text: 待检测文本。

    Returns:
        int: ``str.isalnum`` 为真的字符数。
    """
    return sum(1 for ch in text if ch.isalnum())


def _ensure_not_scan_file(full_text: str) -> None:
    """空文本检测红线：有效字符 < MIN_VALID_CHARS 时抛扫描件信号。

    Args:
        full_text: 解析后的全文文本。

    Raises:
        ScanFileError: 有效字符少于 50（判定为扫描件/图片型文件）。
    """
    if _effective_char_count(full_text) < MIN_VALID_CHARS:
        raise ScanFileError()


# ---------------------------------------------------------------------------
# 标题区间累积器（md / docx 共用）
# ---------------------------------------------------------------------------


class _SectionBuilder:
    """按标题层级流式累积 Section（标题闭合正文区间）。

    标题栈维护当前层级链：遇到 level N 的标题时弹出所有 level >= N
    的旧标题再压入，heading_path 为栈内标题的 `` > `` 连接。
    """

    def __init__(self) -> None:
        """初始化空的标题栈与正文缓冲。"""
        self._stack: list[tuple[int, str]] = []
        self._buf: list[str] = []
        self.sections: list[Section] = []

    @property
    def heading_path(self) -> str:
        """当前标题层级路径。

        Returns:
            str: 栈内标题以 `` > `` 连接；空栈返回空串。
        """
        return " > ".join(title for _, title in self._stack)

    def push_heading(self, level: int, title: str) -> None:
        """压入新标题并闭合当前正文区间。

        Args:
            level: 标题层级（1~6，数字越小层级越高）。
            title: 标题文本。
        """
        self._close()
        while self._stack and self._stack[-1][0] >= level:
            self._stack.pop()
        self._stack.append((level, title))

    def append_text(self, text: str) -> None:
        """向当前区间追加一行正文（空行原样保留）。

        空行是段落分隔符，切分阶段（rag/splitter.py）依赖空行识别
        段落边界做「组内按段落聚合」，故必须保留而非丢弃。

        Args:
            text: 一行正文（可为空行）。
        """
        self._buf.append(text)

    def _close(self) -> None:
        """闭合当前正文区间（缓冲非空才生成 Section）。"""
        body = "\n".join(self._buf).strip("\n")
        if body.strip():
            self.sections.append(
                Section(heading_path=self.heading_path, text=body)
            )
        self._buf = []

    def finish(self) -> list[Section]:
        """闭合最后的正文区间并返回全部 Section。

        Returns:
            list[Section]: 标题区间列表。
        """
        self._close()
        return self.sections


# ---------------------------------------------------------------------------
# 各格式解析器
# ---------------------------------------------------------------------------


def _load_plain_text(path: Path, *, is_markdown: bool) -> ParsedDocument:
    """解析纯文本类文档（MD / TXT）。

    md 以行级 ``#`` 正则识别标题层级；txt 无标题结构（单区间）。

    Args:
        path: 文件路径。
        is_markdown: True 按 md 标题规则解析，False 按 txt 原文解析。

    Returns:
        ParsedDocument: 解析结果。

    Raises:
        ValueError: 编码无法识别。
    """
    full_text = _decode_text_bytes(path.read_bytes())
    builder = _SectionBuilder()
    for line in full_text.splitlines():
        if is_markdown:
            matched = _MD_HEADING_RE.match(line)
            if matched:
                builder.push_heading(len(matched.group(1)), matched.group(2))
                continue
        builder.append_text(line)
    sections = builder.finish() or [Section(heading_path="", text=full_text.strip())]
    return ParsedDocument(sections=sections, full_text=full_text)


def _load_pdf(path: Path) -> ParsedDocument:
    """解析 PDF：pypdf 逐页抽取文本（每页一个 Section）。

    Args:
        path: 文件路径。

    Returns:
        ParsedDocument: 解析结果（PDF 无标题结构，heading_path 为空串）。

    Raises:
        RuntimeError: pypdf 未安装（防御性导入，按「解析失败」口径）。
        Exception: PDF 结构损坏等解析异常（由流水线转 failed）。
    """
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - 环境缺依赖的防御分支
        raise RuntimeError("解析库 pypdf 不可用") from exc
    reader = PdfReader(str(path))
    sections: list[Section] = []
    page_texts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        page_texts.append(text)
        if text.strip():
            sections.append(Section(heading_path="", text=text.strip()))
    full_text = "\n".join(page_texts)
    if not sections:
        sections = [Section(heading_path="", text="")]
    return ParsedDocument(sections=sections, full_text=full_text)


def _load_docx(path: Path) -> ParsedDocument:
    """解析 DOCX：段落 + 表格文本，Heading 样式切标题区间。

    Args:
        path: 文件路径。

    Returns:
        ParsedDocument: 解析结果。

    Raises:
        RuntimeError: python-docx 未安装（防御性导入）。
    """
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("解析库 python-docx 不可用") from exc
    document = Document(str(path))
    builder = _SectionBuilder()
    for paragraph in document.paragraphs:
        style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
        text = paragraph.text.strip()
        if not text:
            continue
        if style_name.startswith("Heading") or style_name.startswith("标题"):
            level = _docx_heading_level(style_name)
            builder.push_heading(level, text)
        else:
            builder.append_text(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                builder.append_text("\t".join(cells))
    sections = builder.finish()
    full_text = "\n".join(section["text"] for section in sections)
    return ParsedDocument(sections=sections, full_text=full_text)


def _docx_heading_level(style_name: str) -> int:
    """从 docx 样式名提取标题层级（Heading 1 / 标题 2 → 1 / 2）。

    Args:
        style_name: 样式名。

    Returns:
        int: 层级数字；无法解析时默认 2（正文次级标题）。
    """
    digits = re.search(r"(\d+)", style_name)
    if digits:
        return max(1, min(6, int(digits.group(1))))
    return 2


def _load_pptx(path: Path) -> ParsedDocument:
    """解析 PPTX：按 slide 抽 text_frame（slide 标题 → 标题节点）。

    Args:
        path: 文件路径。

    Returns:
        ParsedDocument: 解析结果（每个含文本的 slide 一个 Section，
        heading_path 为 slide 标题）。

    Raises:
        RuntimeError: python-pptx 未安装（防御性导入）。
    """
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("解析库 python-pptx 不可用") from exc
    presentation = Presentation(str(path))
    sections: list[Section] = []
    slide_texts: list[str] = []
    for slide in presentation.slides:
        slide_title = ""
        body_lines: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            frame_text = shape.text_frame.text.strip()
            if not frame_text:
                continue
            if shape == slide.shapes.title and not slide_title:
                slide_title = frame_text
            else:
                body_lines.append(frame_text)
        slide_texts.append("\n".join([slide_title, *body_lines]))
        if body_lines or slide_title:
            sections.append(
                Section(heading_path=slide_title, text="\n".join(body_lines))
            )
    full_text = "\n".join(slide_texts)
    if not sections:
        sections = [Section(heading_path="", text="")]
    return ParsedDocument(sections=sections, full_text=full_text)


# ---------------------------------------------------------------------------
# 对外入口
# ---------------------------------------------------------------------------

_PLAIN_SUFFIXES: frozenset[str] = frozenset({"md", "txt", "markdown"})


def load_document(path: Path, source_type: str) -> ParsedDocument:
    """按扩展名路由解析器并返回结构化解析结果（含空文本检测）。

    Args:
        path: 文件绝对路径。
        source_type: 文档类型（pdf/md/txt/docx/pptx，来自
            knowledge_items.source_type；实际路由以扩展名为准，
            两者不一致时以扩展名为权威——文件本体决定解析方式）。

    Returns:
        ParsedDocument: sections（标题区间列表）+ full_text。

    Raises:
        ScanFileError: 解析后有效字符 < 50（扫描件/图片型文件红线，
            文案逐字锁定为 SCAN_FILE_ERROR_MSG）。
        ValueError: TXT 编码无法识别。
        RuntimeError: 对应格式解析库未安装（防御性导入口径）。
        Exception: 解析库内部异常（截断 PDF 等，由流水线转
            failed「解析失败：{摘要}」）。
    """
    suffix = path.suffix.lower().lstrip(".")
    if suffix in _PLAIN_SUFFIXES:
        parsed = _load_plain_text(path, is_markdown=suffix != "txt")
    elif suffix == "pdf":
        parsed = _load_pdf(path)
    elif suffix == "docx":
        parsed = _load_docx(path)
    elif suffix == "pptx":
        parsed = _load_pptx(path)
    else:
        # 受理层已做白名单拦截；此处兜底不可达分支，按解析失败口径抛出。
        raise ValueError(f"不支持的文件扩展名: .{suffix}")
    _ensure_not_scan_file(parsed["full_text"])
    return parsed