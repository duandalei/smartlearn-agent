"""loader 解析逻辑单元测试（tasks.md 5.3）。

覆盖：空文本红线（文案逐字锁定）/ 阈值边界 / TXT 编码回退链（UTF-8/
GBK/UTF-16 与全失败）/ md 多级标题栈 / PDF、DOCX、PPTX 最小真实样例
/ 损坏字节流异常口径。PDF 与 Office 样例均为测试内动态生成的最小合法
文件（依赖已安装的 pypdf / python-docx / python-pptx）。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from app.rag.loader import (
    MIN_VALID_CHARS,
    SCAN_FILE_ERROR_MSG,
    ScanFileError,
    load_document,
)


def _write(tmp_path: Path, name: str, data: bytes | str) -> Path:
    """写入临时文件并返回路径（str 按 UTF-8 编码）。"""
    path = tmp_path / name
    path.write_bytes(data if isinstance(data, bytes) else data.encode("utf-8"))
    return path


def _make_pdf(lines: list[str]) -> bytes:
    """构造最小合法单页 PDF（Helvetica 文本，动态计算 xref 偏移）。"""
    parts = ["BT", "/F1 10 Tf", "72 720 Td", "14 TL"]
    for line in lines:
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        parts.append(f"({escaped}) Tj T*")
    parts.append("ET")
    content = "\n".join(parts).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for idx, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{idx} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 6\n0000000000 65535 f \n"
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += (
        b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n"
        + str(xref_pos).encode()
        + b"\n%%EOF"
    )
    return bytes(out)


# ---------------------------------------------------------------------------
# 空文本检测红线
# ---------------------------------------------------------------------------


def test_scan_file_redline_exact_message(tmp_path: Path) -> None:
    """全标点/空白内容触发红线，文案逐字锁定（spec 5.2.1 规则 2）。"""
    path = _write(tmp_path, "scan.txt", "！" * 200 + "\n\n……" * 50)
    with pytest.raises(ScanFileError) as exc_info:
        load_document(path, "txt")
    assert str(exc_info.value) == SCAN_FILE_ERROR_MSG
    assert SCAN_FILE_ERROR_MSG == "该文件可能是扫描件，未提取到文字（当前版本不支持 OCR）"


def test_scan_file_threshold_boundary(tmp_path: Path) -> None:
    """阈值边界：49 个有效字符触发红线，50 个恰好通过。"""
    path49 = _write(tmp_path, "a.txt", "x" * (MIN_VALID_CHARS - 1) + " " * 10)
    with pytest.raises(ScanFileError):
        load_document(path49, "txt")
    path50 = _write(tmp_path, "b.txt", "x" * MIN_VALID_CHARS + " ")
    doc = load_document(path50, "txt")
    assert doc["full_text"].strip() == "x" * MIN_VALID_CHARS


# ---------------------------------------------------------------------------
# TXT 编码回退链
# ---------------------------------------------------------------------------


def test_encoding_fallback_chain(tmp_path: Path) -> None:
    """UTF-8 / GBK / UTF-16(含 BOM) 三种编码文件均解出原文。"""
    text = (
        "机器学习是人工智能的核心分支，监督学习与无监督学习是其两大范式，"
        "深度学习则依托多层神经网络自动学习数据的层次化表示。"
    )
    for name, encoding in (("u8.txt", "utf-8"), ("gk.txt", "gbk"), ("u16.txt", "utf-16")):
        path = _write(tmp_path, name, text.encode(encoding))
        doc = load_document(path, "txt")
        assert text in doc["full_text"], f"{encoding} 编码解码失败"


def test_encoding_all_failed_raises_value_error(tmp_path: Path) -> None:
    """三链全失败（孤立代理字节）抛 ValueError。"""
    path = _write(tmp_path, "bad.txt", b"\xff\xfe\x00\xd8\x00\xd8")
    with pytest.raises(ValueError, match="编码无法识别"):
        load_document(path, "txt")


# ---------------------------------------------------------------------------
# Markdown 标题栈
# ---------------------------------------------------------------------------


def test_markdown_heading_stack_levels(tmp_path: Path) -> None:
    """多级标题栈：## 弹出 ###、保留 #；heading_path 以 ' > ' 连接。"""
    md = "\n".join(
        [
            "# 一级总览",
            "总览正文内容，介绍整体结构与阅读脉络，帮助读者建立全局认识。",
            "## 二级甲",
            "甲下正文内容，展开第一主题的关键概念与典型应用场景供参考。",
            "### 三级细则",
            "细则正文内容，给出公式推导与代码示例，方便动手实践验证。",
            "## 二级乙",
            "乙下正文内容，讨论第二主题的进阶话题与常见误区提醒。",
        ]
    )
    doc = load_document(_write(tmp_path, "h.md", md), "md")
    paths = [s["heading_path"] for s in doc["sections"]]
    assert paths == ["一级总览", "一级总览 > 二级甲", "一级总览 > 二级甲 > 三级细则", "一级总览 > 二级乙"]
    assert doc["sections"][2]["text"].strip() == "细则正文内容，给出公式推导与代码示例，方便动手实践验证。"


# ---------------------------------------------------------------------------
# PDF / DOCX / PPTX 最小真实样例
# ---------------------------------------------------------------------------


def test_pdf_minimal_sample(tmp_path: Path) -> None:
    """最小 PDF：文本可提取（英文 ASCII），无标题结构。"""
    lines = [
        "Machine learning is a subfield of artificial intelligence.",
        "It focuses on algorithms that learn patterns from data.",
    ]
    path = _write(tmp_path, "mini.pdf", _make_pdf(lines))
    doc = load_document(path, "pdf")
    full = doc["full_text"]
    assert "Machine learning" in full and "algorithms" in full
    assert all(s["heading_path"] == "" for s in doc["sections"])


def test_docx_minimal_sample(tmp_path: Path) -> None:
    """最小 DOCX：Heading 样式切区间，表格文本并入正文。"""
    document = Document()
    document.add_heading("深度学习基础", level=1)
    document.add_paragraph("深度学习使用多层神经网络自动学习数据表示，广泛用于视觉与语音。")
    document.add_heading("训练技巧", level=2)
    document.add_paragraph("学习率调度与梯度裁剪是稳定训练的常用技巧，不可忽视。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "优化器"
    table.cell(0, 1).text = "Adam"
    path = tmp_path / "mini.docx"
    document.save(str(path))
    doc = load_document(path, "docx")
    paths = [s["heading_path"] for s in doc["sections"]]
    assert "深度学习基础" in paths
    assert "深度学习基础 > 训练技巧" in paths
    assert "优化器\tAdam" in doc["full_text"]


def test_pptx_minimal_sample(tmp_path: Path) -> None:
    """最小 PPTX：slide 标题作 heading_path，正文入 text。"""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "注意力机制"
    body = slide.placeholders[1]
    body.text = (
        "注意力机制让模型动态聚焦相关信息，是 Transformer 架构的核心组件，"
        "使序列建模摆脱循环依赖并支持并行训练。"
    )
    path = tmp_path / "mini.pptx"
    presentation.save(str(path))
    doc = load_document(path, "pptx")
    assert len(doc["sections"]) == 1
    assert doc["sections"][0]["heading_path"] == "注意力机制"
    assert "Transformer" in doc["sections"][0]["text"]


def test_corrupted_pdf_raises_not_scan(tmp_path: Path) -> None:
    """损坏字节流：抛解析异常而非扫描件红线（流水线按解析失败处理）。"""
    path = _write(tmp_path, "broken.pdf", b"this is definitely not a pdf file 1234567890")
    with pytest.raises(Exception) as exc_info:
        load_document(path, "pdf")
    assert not isinstance(exc_info.value, ScanFileError)